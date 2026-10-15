import sys

import json
import socket
import time
import tempfile
import shutil
from datetime import datetime, timezone
from pathlib import Path
import tkinter as tk
from tkinter import filedialog

import pandas as pd
from flask import Flask, jsonify, render_template, request
from flask_cors import CORS
from openpyxl import load_workbook

IS_FROZEN = getattr(sys, "frozen", False)
BUNDLE_DIR = Path(sys._MEIPASS) if IS_FROZEN else Path(__file__).resolve().parent
APP_DIR = Path(sys.executable).resolve().parent if IS_FROZEN else Path(__file__).resolve().parent
BASE_DIR = APP_DIR
UPLOAD_DIR = BASE_DIR / "uploads"
SAMPLE_FILE = BUNDLE_DIR / "sample_data.xlsx"

UPLOAD_DIR.mkdir(exist_ok=True)

app = Flask(
    __name__,
    template_folder=str(BUNDLE_DIR / "templates"),
    static_folder=str(BUNDLE_DIR / "static"),
)
CORS(app)

CURRENT_FILE: Path | None = None
FILTER_COLUMNS = (
    "Year",
    "Name of Student",
    "Gender",
    "Affiliation",
    "State",
    "Branch",
    "Branch Category",
    "Field / Branch of Engg.",
    "Division (DRDO)",
    "Scientist Name",
    "Designation",
    "Mode of Training",
    "Certificate Issued",
)
EXCEL_COLUMNS = (
    "Year",
    "S.No.",
    "Name of Student",
    "Gender",
    "Affiliation",
    "State",
    "Branch",
    "Branch Category",
    "Field / Branch of Engg.",
    "Division (DRDO)",
    "Scientist Name",
    "Designation",
    "Period of Attachment",
    "Time in Weeks",
    "Mode of Training",
    "Training Period",
    "Topic Name",
    "Certificate Issued",
    "Certificate Issue Date",
)
FORM_REQUIRED_FIELDS = (
    "Year",
    "Name of Student",
    "Gender",
    "Affiliation",
    "State",
    "Branch",
    "Branch Category",
    "Field / Branch of Engg.",
    "Division (DRDO)",
    "Scientist Name",
    "Designation",
    "Period of Attachment",
    "Time in Weeks",
    "Mode of Training",
    "Training Period",
)
DIVISION_OPTIONS = (
    "Laser Technology",
    "MEMS",
    "THZ & NANO",
    "III-V Technologies",
    "Cryocoolars",
    "Laser System Engineering Group",
    "Magnetometer",
    "Atomic Clock",
    "Instrumentation",
    "Futuristic Technologies",
    "HRD",
    "TIRC",
    "THz & Quantum Measurements",
    "Assembly & RF Characterization",
    "Device Modelling and PDK Development",
    "MMIC Design",
    "IR Sensor & Array Evaluation",
    "Saw based Systems",
    "Photonics Integrated Circuits",
    "ROIC & PROXIMITY DESIGN",
    "Acoustic Systems",
    "CHARACTERIZATION",
    "Silicon Carbide Growth",
    "CZT/PMN-PT",
    "LPE",
    "MOCVD",
    "MBE",
    "Ion Implantation",
)
SYNTHETIC_NUMERIC_COLUMNS = ("Count",)
RESERVED_QUERY_KEYS = {"page", "per_page", "x_axis", "y_axis", "filter_col", "filter_val"}

SCIENTISTS_CSV = BASE_DIR / "Scientists_Name.csv"
DROPDOWN_OPTIONS_FILE = BASE_DIR / "dropdown_options.json"

BRANCH_OPTIONS = (
    "Computer Science Engineering",
    "Information Technology",
    "Electronics & Communication Engineering",
    "Electrical Engineering",
    "Electrical & Electronics Engineering",
    "Mechanical Engineering",
    "Civil Engineering",
    "Chemical Engineering",
    "Aerospace Engineering",
    "Instrumentation Engineering",
    "Biotechnology Engineering",
    "Biomedical Engineering",
    "Materials Science & Engineering",
    "Production Engineering",
    "Industrial Engineering",
    "Metallurgical Engineering",
    "Environmental Engineering",
    "Agriculture Engineering",
    "Mining Engineering",
    "Textile Engineering",
)
BRANCH_CATEGORY_OPTIONS = (
    "B.Tech",
    "B.E.",
    "B.Tech (Hons.)",
    "M.Tech",
    "M.E.",
    "M.Tech (Integrated)",
    "Dual Degree (B.Tech+M.Tech)",
)
OPTION_FIELDS = {"branch", "branch_category", "division", "scientist", "designation"}


def _load_scientists_csv() -> tuple[list[str], list[str]]:
    """
    Load scientist names and designations from the official Scientists_Name.csv.
    Returns (names_list, designations_list) sorted and deduplicated.
    Works offline - file is local to the project.
    """
    import csv
    names: list[str] = []
    desigs: set[str] = set()
    try:
        with open(str(SCIENTISTS_CSV), "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                raw_title = (row.get("Title") or "").strip().rstrip(".")
                raw_name = (row.get("Name") or "").strip()
                raw_desig = (row.get("Designation") or "").strip()
                if raw_name:
                    # Format: "TITLE. NAME" - title period is added by us, not from CSV
                    full_name = f"{raw_title}. {raw_name}" if raw_title else raw_name
                    names.append(full_name)
                if raw_desig:
                    desigs.add(raw_desig)
    except (FileNotFoundError, IOError):
        pass
    return sorted(set(names)), sorted(desigs)


def _lan_ip() -> str:
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"


def _read_excel(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return pd.read_excel(path, engine="openpyxl")
    if suffix == ".xls":
        try:
            import xlrd  # noqa: F401
        except ImportError:
            raise ValueError("Reading .xls files requires xlrd. Install it with: pip install xlrd") from None
        return pd.read_excel(path, engine="xlrd")
    raise ValueError("Unsupported file type. Use .xlsx or .xls")


def _active_path() -> Path:
    global CURRENT_FILE
    if CURRENT_FILE and CURRENT_FILE.exists():
        return CURRENT_FILE
    if SAMPLE_FILE.exists():
        CURRENT_FILE = SAMPLE_FILE
        return SAMPLE_FILE
    raise FileNotFoundError("No data file loaded. Upload an Excel file.")


def _file_info(path: Path) -> dict:
    stat = path.stat()
    modified = datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat()
    return {
        "filename": path.name,
        "filepath": str(path.resolve()),
        "file_size": int(stat.st_size),
        "modified_at_utc": modified,
    }


def _load_df() -> pd.DataFrame:
    df = _read_excel(_active_path())
    df = df.where(pd.notnull(df), None)
    return df


MONTH_ORDER = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}


def _normalize(text: str) -> str:
    return "".join(ch.lower() for ch in str(text) if ch.isalnum())


def _normalize_field(text: str) -> str:
    """
    Normalize a field value for consistent comparison.
    - Removes all punctuation/dashes/special chars
    - Removes extra spaces
    - Case-insensitive
    - Keeps only alphanumeric characters and single spaces
    """
    if not text:
        return ""
    text = str(text).strip().lower()
    # Replace non-alphanumeric with spaces to preserve word boundaries
    text = "".join(ch if ch.isalnum() else " " for ch in text)
    # Collapse multiple spaces
    text = " ".join(text.split())
    return text


def _duplicate_key(row: dict) -> tuple[str, str, str, str]:
    """
    Generate a normalized key for duplicate detection.
    Compares: Name + Affiliation + Division + Period (all normalized).
    """
    return (
        _normalize_field(row.get("Name of Student", "")),
        _normalize_field(row.get("Affiliation", "")),
        _normalize_field(row.get("Division (DRDO)", "")),
        _normalize_field(row.get("Period of Attachment", "")),
    )


def _pick_column(df: pd.DataFrame, aliases: tuple[str, ...], numeric_only: bool = False) -> str | None:
    alias_set = {_normalize(alias) for alias in aliases}
    for col in df.columns:
        if _normalize(col) in alias_set:
            if not numeric_only:
                return str(col)
            series = pd.to_numeric(df[col], errors="coerce")
            if series.notna().any():
                return str(col)
    return None


def _sort_chart_labels(df: pd.DataFrame, x_axis: str) -> pd.DataFrame:
    if pd.api.types.is_datetime64_any_dtype(df[x_axis]):
        return df.sort_values(x_axis)

    if pd.api.types.is_numeric_dtype(df[x_axis]):
        return df.sort_values(x_axis)

    if df[x_axis].dtype == object or pd.api.types.is_string_dtype(df[x_axis]):
        candidate = df[x_axis].dropna().astype(str).str.strip().str.lower()
        if candidate.isin(MONTH_ORDER).any():
            order = candidate.map(lambda v: MONTH_ORDER.get(v, 99))
            return df.assign(_month_order=order).sort_values(["_month_order", x_axis]).drop(columns=["_month_order"])

    return df


def _apply_filters(df: pd.DataFrame, args) -> pd.DataFrame:
    for col, val in args.items():
        if col in RESERVED_QUERY_KEYS:
            continue
        if val and str(val).lower() != "all" and col in df.columns:
            series = df[col]
            if pd.api.types.is_numeric_dtype(series):
                try:
                    df = df[series == float(val) if "." in str(val) else series == int(val)]
                except (ValueError, TypeError):
                    df = df[series.astype(str) == str(val)]
            else:
                df = df[series.astype(str) == str(val)]
    return df


def _df_to_records(df: pd.DataFrame) -> list[dict]:
    records = df.to_dict(orient="records")
    for row in records:
        for key, value in row.items():
            if pd.isna(value):
                row[key] = None
            elif hasattr(value, "item"):
                try:
                    row[key] = value.item()
                except (ValueError, AttributeError):
                    row[key] = value
    return records


def _clean_form_value(value) -> str:
    """Clean and trim form input value."""
    return "" if value is None else str(value).strip()


def _round_whole_number(value, field_name: str) -> int:
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        raise ValueError(f"{field_name} must be a number") from None


def _normalize_form_value(value: str, normalize: bool = True) -> str:
    """
    Normalize form input values for storage in Excel.
    By default, applies normalization (lowercase, remove extra punctuation).
    Set normalize=False to just clean/trim without changing case.
    """
    cleaned = _clean_form_value(value)
    if normalize and cleaned:
        return _normalize_field(cleaned)
    return cleaned


def _pick_column(df: pd.DataFrame, aliases: tuple[str, ...], numeric_only: bool = False) -> str | None:
    alias_set = {_normalize(alias) for alias in aliases}
    for col in df.columns:
        if _normalize(col) in alias_set:
            if not numeric_only:
                return str(col)
            series = pd.to_numeric(df[col], errors="coerce")
            if series.notna().any():
                return str(col)
    return None


def _workbook_headers(sheet) -> list[str]:
    """Extract headers from Excel sheet (first row)."""
    return [_clean_form_value(cell.value) for cell in sheet[1]]


def _save_workbook_safe(workbook, path: Path, max_retries: int = 3, retry_delay: float = 0.5):
    """
    Attempt to save workbook with retry logic for locked files.
    On Windows, files locked by Excel are often freed after a short delay.
    """
    last_error = None
    for attempt in range(max_retries):
        try:
            workbook.save(path)
            return  # Success
        except PermissionError as e:
            last_error = e
            if attempt < max_retries - 1:
                time.sleep(retry_delay)
                continue
        except Exception as e:
            last_error = e
            break
    
    if last_error:
        raise last_error


def _scientist_options() -> list[str]:
    """Return sorted list of scientist names from the official CSV file."""
    names, _ = _load_scientists_csv()
    return names


def _designation_options() -> list[str]:
    """Return sorted list of unique designations from the official CSV file."""
    _, designations = _load_scientists_csv()
    return designations


def _option_key(value: str) -> str:
    return _normalize_field(value)


def _dedupe_options(values: list[str] | tuple[str, ...]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = _clean_form_value(value)
        key = _option_key(cleaned)
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
    return result


def _default_dropdown_options() -> dict[str, list[str]]:
    return {
        "branch": _dedupe_options(BRANCH_OPTIONS),
        "branch_category": _dedupe_options(BRANCH_CATEGORY_OPTIONS),
        "division": _dedupe_options(DIVISION_OPTIONS),
        "scientist": _dedupe_options(_scientist_options() + ["Other"]),
        "designation": _dedupe_options(_designation_options()),
    }


def _empty_dropdown_state() -> dict[str, dict[str, list[str]]]:
    return {field: {"added": [], "removed": []} for field in OPTION_FIELDS}


def _load_dropdown_state() -> dict[str, dict[str, list[str]]]:
    state = _empty_dropdown_state()
    if not DROPDOWN_OPTIONS_FILE.exists():
        return state
    try:
        raw = json.loads(DROPDOWN_OPTIONS_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return state
    if not isinstance(raw, dict):
        return state
    for field in OPTION_FIELDS:
        field_state = raw.get(field, {})
        if not isinstance(field_state, dict):
            continue
        for bucket in ("added", "removed"):
            values = field_state.get(bucket, [])
            if isinstance(values, list):
                state[field][bucket] = _dedupe_options(values)
    return state


def _save_dropdown_state(state: dict[str, dict[str, list[str]]]):
    payload = {
        field: {
            "added": _dedupe_options(state.get(field, {}).get("added", [])),
            "removed": _dedupe_options(state.get(field, {}).get("removed", [])),
        }
        for field in sorted(OPTION_FIELDS)
    }
    tmp = DROPDOWN_OPTIONS_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    tmp.replace(DROPDOWN_OPTIONS_FILE)


def _merged_dropdown_options() -> dict[str, list[str]]:
    defaults = _default_dropdown_options()
    state = _load_dropdown_state()
    merged: dict[str, list[str]] = {}
    for field, default_values in defaults.items():
        removed = {_option_key(value) for value in state[field]["removed"]}
        values = [value for value in default_values if _option_key(value) not in removed]
        for value in state[field]["added"]:
            key = _option_key(value)
            if key not in removed and all(_option_key(existing) != key for existing in values):
                values.append(value)
        merged[field] = values
    return merged


def _update_dropdown_option(action: str, field: str, value: str) -> list[str]:
    field = _clean_form_value(field)
    value = _clean_form_value(value)
    if field not in OPTION_FIELDS:
        raise ValueError("Invalid dropdown field")
    if not value:
        raise ValueError("Option value is required")

    state = _load_dropdown_state()
    defaults = _default_dropdown_options()
    value_key = _option_key(value)

    state[field]["added"] = [
        existing for existing in state[field]["added"] if _option_key(existing) != value_key
    ]
    state[field]["removed"] = [
        existing for existing in state[field]["removed"] if _option_key(existing) != value_key
    ]

    if action == "add":
        if all(_option_key(existing) != value_key for existing in defaults[field]):
            state[field]["added"].append(value)
    elif action == "remove":
        state[field]["removed"].append(value)
    else:
        raise ValueError("Invalid dropdown action")

    _save_dropdown_state(state)
    return _merged_dropdown_options()[field]


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/form")
def entry_form():
    dropdown_options = _merged_dropdown_options()
    return render_template(
        "form.html",
        current_year=datetime.now().year,
        dropdown_options=dropdown_options,
        division_options=dropdown_options["division"],
        scientist_options=dropdown_options["scientist"],
        designation_options=dropdown_options["designation"],
    )


@app.route("/dropdown-options")
def dropdown_options():
    return jsonify(_merged_dropdown_options())


@app.route("/dropdown-option", methods=["POST"])
def dropdown_option():
    payload = request.get_json(silent=True) or {}
    try:
        options = _update_dropdown_option(
            _clean_form_value(payload.get("action")),
            _clean_form_value(payload.get("field")),
            _clean_form_value(payload.get("value")),
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except OSError as exc:
        return jsonify({"error": f"Unable to save dropdown option: {exc}"}), 500
    return jsonify({"status": "ok", "options": options})


@app.route("/submit-form", methods=["POST"])
def submit_form():
    payload = request.get_json(silent=True) or {}
    data = {column: _clean_form_value(payload.get(column)) for column in EXCEL_COLUMNS}
    data["Year"] = data["Year"] or str(datetime.now().year)

    missing = [field for field in FORM_REQUIRED_FIELDS if not data.get(field)]
    if missing:
        return jsonify({"error": f"Missing required fields: {', '.join(missing)}"}), 400

    try:
        data["Time in Weeks"] = _round_whole_number(data["Time in Weeks"], "Time in Weeks")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    try:
        data["Year"] = int(data["Year"])
    except (TypeError, ValueError):
        data["Year"] = datetime.now().year

    try:
        path = _active_path()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    if path.suffix.lower() != ".xlsx":
        return jsonify({"error": "Form submissions can only be saved to .xlsx files"}), 400

    try:
        workbook = load_workbook(path)
        sheet = workbook.active
        headers = _workbook_headers(sheet)
        header_map = {header: index + 1 for index, header in enumerate(headers) if header}
        missing_columns = [column for column in EXCEL_COLUMNS if column not in header_map]
        if missing_columns:
            return jsonify({"error": f"Excel file is missing columns: {', '.join(missing_columns)}"}), 400

        incoming_key = _duplicate_key(data)
        sno_values = []
        sno_col_idx = header_map.get("S.No.")
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = {
                header: row[index - 1] if index - 1 < len(row) else None
                for header, index in header_map.items()
            }
            if _duplicate_key(row_data) == incoming_key:
                workbook.close()
                return jsonify({"error": "Duplicate entry detected. This student record already exists."}), 409
            if sno_col_idx and sno_col_idx - 1 < len(row):
                sno_val = row[sno_col_idx - 1]
                if sno_val is not None:
                    try:
                        sno_values.append(int(sno_val))
                    except (TypeError, ValueError):
                        continue

        if sno_values:
            new_sno = max(sno_values) + 1
        else:
            new_sno = 1
        data["S.No."] = new_sno

        next_row = sheet.max_row + 1
        for column in EXCEL_COLUMNS:
            sheet.cell(row=next_row, column=header_map[column], value=data.get(column))
        
        # Use safe save with retry logic for locked files
        _save_workbook_safe(workbook, path, max_retries=3, retry_delay=0.5)
        workbook.close()
        
    except PermissionError:
        return jsonify({"error": "Excel file is open or locked. Close it in Excel and try again. (Automatic retry in 1-2 seconds may help)"}), 423
    except Exception as exc:
        return jsonify({"error": f"Unable to save entry: {str(exc)}"}), 500

    return jsonify({"message": "Entry submitted successfully", "sno": new_sno}), 201


@app.route("/get-entry", methods=["GET"])
def get_entry():
    sno_param = request.args.get("sno")
    name_param = request.args.get("name")

    if not sno_param and not name_param:
        return jsonify({"error": "Please provide at least S.No. or Name of Student to search"}), 400

    try:
        df = _load_df()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    mask = pd.Series([True] * len(df))

    if sno_param:
        try:
            sno = int(sno_param)
        except (TypeError, ValueError):
            return jsonify({"error": "sno must be a number"}), 400
        if "S.No." not in df.columns:
            return jsonify({"error": "Excel file is missing S.No. column"}), 400
        mask = mask & (df["S.No."] == sno)

    if name_param:
        if "Name of Student" not in df.columns:
            return jsonify({"error": "Excel file is missing Name of Student column"}), 400
        name_clean = name_param.strip()
        mask = mask & (df["Name of Student"].astype(str).str.strip().str.lower() == name_clean.lower())

    matched_rows = df[mask]
    if matched_rows.empty:
        return jsonify({"error": "Entry not found"}), 404

    # If searching by name only (no sno) and multiple matches found,
    # return a list of matches for the user to choose from.
    if not sno_param and name_param and len(matched_rows) > 1:
        identify_fields = ["S.No.", "Name of Student", "Affiliation", "Period of Attachment", "Division (DRDO)", "Year"]
        matches = []
        for _, row in matched_rows.iterrows():
            match_obj = {}
            for field in identify_fields:
                if field in row.index:
                    val = row[field]
                    match_obj[field] = str(val) if pd.notna(val) else ""
                else:
                    match_obj[field] = ""
            matches.append(match_obj)
        return jsonify({"multiple": True, "matches": matches})

    row = matched_rows.iloc[0]
    result = {col: str(row[col]) if pd.notna(row[col]) else "" for col in row.index}
    return jsonify(result)


@app.route("/update-entry", methods=["POST"])
def update_entry():
    payload = request.get_json(silent=True) or {}
    original_sno = payload.get("original_sno")

    if not original_sno:
        return jsonify({"error": "original_sno is required"}), 400

    try:
        original_sno = int(original_sno)
    except (TypeError, ValueError):
        return jsonify({"error": "original_sno must be a number"}), 400

    data = {column: _clean_form_value(payload.get(column)) for column in EXCEL_COLUMNS}

    missing = [field for field in FORM_REQUIRED_FIELDS if not data.get(field)]
    if missing:
        return jsonify({"error": f"Missing required fields: {', '.join(missing)}"}), 400

    try:
        data["Time in Weeks"] = _round_whole_number(data["Time in Weeks"], "Time in Weeks")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    try:
        data["Year"] = int(data["Year"])
    except (TypeError, ValueError):
        data["Year"] = datetime.now().year

    try:
        path = _active_path()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    if path.suffix.lower() != ".xlsx":
        return jsonify({"error": "Form submissions can only be saved to .xlsx files"}), 400

    try:
        workbook = load_workbook(path)
        sheet = workbook.active
        headers = _workbook_headers(sheet)
        header_map = {header: index + 1 for index, header in enumerate(headers) if header}
        missing_columns = [column for column in EXCEL_COLUMNS if column not in header_map]
        if missing_columns:
            workbook.close()
            return jsonify({"error": f"Excel file is missing columns: {', '.join(missing_columns)}"}), 400

        if "S.No." not in header_map:
            workbook.close()
            return jsonify({"error": "Excel file is missing S.No. column"}), 400

        target_row = None
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            sno_col_idx = header_map.get("S.No.")
            if sno_col_idx and sno_col_idx - 1 < len(row):
                try:
                    row_sno = int(row[sno_col_idx - 1])
                    if row_sno == original_sno:
                        target_row = row_idx
                        break
                except (TypeError, ValueError):
                    continue

        if target_row is None:
            workbook.close()
            return jsonify({"error": f"No entry found with S.No. {original_sno}"}), 404

        for column in EXCEL_COLUMNS:
            if column == "S.No.":
                continue
            col_idx = header_map.get(column)
            if col_idx:
                sheet.cell(row=target_row, column=col_idx, value=data.get(column))

        _save_workbook_safe(workbook, path, max_retries=3, retry_delay=0.5)
        workbook.close()

    except PermissionError:
        return jsonify({"error": "Excel file is open or locked. Close it in Excel and try again."}), 423
    except Exception as exc:
        return jsonify({"error": f"Unable to update entry: {str(exc)}"}), 500

    return jsonify({"message": "Entry updated successfully"}), 200


@app.route("/edit-form")
def edit_form():
    dropdown_options = _merged_dropdown_options()
    return render_template(
        "edit_form.html",
        current_year=datetime.now().year,
        dropdown_options=dropdown_options,
        division_options=dropdown_options["division"],
        scientist_options=dropdown_options["scientist"],
        designation_options=dropdown_options["designation"],
    )


@app.route("/upload", methods=["POST"])
def upload():
    global CURRENT_FILE
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename"}), 400

    ext = Path(file.filename).suffix.lower()
    if ext not in (".xlsx", ".xls"):
        return jsonify({"error": "Only .xlsx and .xls files are supported"}), 400

    safe_name = Path(file.filename).name
    dest = UPLOAD_DIR / safe_name
    file.save(dest)
    CURRENT_FILE = dest

    return jsonify(
        {
            "message": "File uploaded successfully",
            "filename": safe_name,
            "filepath": str(dest.resolve()),
        }
    )


@app.route("/set-active-file", methods=["POST"])
def set_active_file():
    global CURRENT_FILE
    payload = request.get_json(silent=True) or {}
    raw_path = payload.get("filepath")
    if not raw_path:
        return jsonify({"error": "filepath is required"}), 400

    path = Path(str(raw_path)).expanduser()
    if not path.is_absolute():
        path = (BASE_DIR / path).resolve()
    if not path.exists() or not path.is_file():
        return jsonify({"error": "File not found"}), 404
    if path.suffix.lower() not in (".xlsx", ".xls"):
        return jsonify({"error": "Only .xlsx and .xls files are supported"}), 400

    CURRENT_FILE = path
    payload = _file_info(path)
    payload["status"] = "ok"
    return jsonify(payload)


@app.route("/pick-file", methods=["POST"])
def pick_file():
    global CURRENT_FILE
    try:
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        selected = filedialog.askopenfilename(
            title="Select Excel file",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        root.destroy()
    except Exception as exc:
        return jsonify({"error": f"File picker unavailable: {exc}"}), 500

    if not selected:
        return jsonify({"error": "No file selected"}), 400

    path = Path(selected)
    if path.suffix.lower() not in (".xlsx", ".xls"):
        return jsonify({"error": "Only .xlsx and .xls files are supported"}), 400
    if not path.exists() or not path.is_file():
        return jsonify({"error": "File not found"}), 404

    CURRENT_FILE = path.resolve()
    payload = _file_info(CURRENT_FILE)
    payload["status"] = "ok"
    return jsonify(payload)


@app.route("/meta")
def meta():
    path = _active_path()
    return jsonify(_file_info(path))


@app.route("/reload")
def reload_data():
    try:
        path = _active_path()
        _read_excel(path)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    payload = _file_info(path)
    payload["status"] = "ok"
    return jsonify(payload)


@app.after_request
def disable_cache(response):
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route("/data")
def data():
    try:
        df = _load_df()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    df = _apply_filters(df, request.args)
    page = max(int(request.args.get("page", 1)), 1)
    per_page = max(int(request.args.get("per_page", 10)), 1)

    total = len(df)
    start = (page - 1) * per_page
    end = start + per_page
    page_df = df.iloc[start:end]

    return jsonify(
        {
            "data": _df_to_records(page_df),
            "total": total,
            "page": page,
            "per_page": per_page,
            "pages": max((total + per_page - 1) // per_page, 1),
        }
    )


@app.route("/kpi")
def kpi():
    try:
        df = _load_df()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    df = _apply_filters(df, request.args)

    student_col = _pick_column(df, ("name of student", "student name", "student", "name"))
    gender_col = _pick_column(df, ("gender", "sex"))
    branch_col = _pick_column(df, ("branch", "branch category", "field / branch of engg.", "field of engineering"))
    institute_col = _pick_column(df, ("affiliation", "institute", "institution", "college", "university"))
    division_col = _pick_column(df, ("division (drdo)", "drdo division", "division", "department"))

    def clean_text(name: str | None) -> pd.Series:
        if not name or name not in df.columns:
            return pd.Series(dtype="object")
        return df[name].dropna().astype(str).str.strip()

    def count_col(name: str | None) -> int:
        if not name or name not in df.columns:
            return int(len(df))
        values = clean_text(name)
        return int(values.ne("").sum())

    def count_gender(target: str) -> int:
        values = clean_text(gender_col).str.lower()
        return int(values.eq(target).sum())

    def unique_col(name: str | None) -> int:
        values = clean_text(name)
        return int(values[values.ne("")].nunique())

    certificate_col = _pick_column(df, ("certificate issued", "certificate"))
    weeks_col = _pick_column(df, ("time in weeks", "weeks", "training weeks", "duration"), numeric_only=True)

    def count_yes(name: str | None) -> int:
        if not name or name not in df.columns:
            return 0
        values = clean_text(name).str.lower()
        return int(values.isin({"yes", "y", "true", "1"}).sum())

    def average_numeric(name: str | None) -> float:
        if not name or name not in df.columns:
            return 0.0
        values = pd.to_numeric(df[name], errors="coerce").dropna()
        return float(values.mean()) if not values.empty else 0.0

    return jsonify(
        {
            "total_students": count_col(student_col),
            "female_interns": count_gender("female"),
            "male_interns": count_gender("male"),
            "branches_involved": unique_col(branch_col),
            "institutes_involved": unique_col(institute_col),
            "drdo_divisions": unique_col(division_col),
            "certificates_issued": count_yes(certificate_col),
            "average_weeks": average_numeric(weeks_col),
        }
    )


@app.route("/columns")
def columns():
    try:
        df = _load_df()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    numeric = [
        c
        for c in df.columns
        if pd.api.types.is_numeric_dtype(df[c]) or pd.to_numeric(df[c], errors="coerce").notna().any()
    ]
    if "Count" not in numeric:
        numeric.append("Count")
    categorical = [c for c in df.columns if c not in numeric or c in FILTER_COLUMNS]

    return jsonify(
        {
            "columns": list(df.columns.astype(str)),
            "numeric": [str(c) for c in numeric],
            "categorical": [str(c) for c in categorical],
        }
    )


@app.route("/filters")
def filters():
    try:
        df = _load_df()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    result = {}
    for col in df.columns:
        if col in FILTER_COLUMNS or df[col].dtype == object or df[col].nunique() <= 25:
            values = df[col].dropna().astype(str).unique().tolist()
            # Normalize for display and remove duplicates caused by formatting
            normalized_map = {}  # Maps normalized value to display value
            for val in values:
                norm_key = _normalize_field(val)
                if norm_key not in normalized_map:
                    normalized_map[norm_key] = val  # Keep first occurrence
            # Use original values but deduplicated
            unique_values = sorted(normalized_map.values())
            result[str(col)] = unique_values
    return jsonify(result)


@app.route("/chart-data")
def chart_data():
    try:
        df = _load_df()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    df = _apply_filters(df, request.args)

    x_axis = request.args.get("x_axis")
    y_axis = request.args.get("y_axis")
    filter_col = request.args.get("filter_col")
    filter_val = request.args.get("filter_val")

    if filter_col and filter_val and filter_col in df.columns:
        if str(filter_val).lower() != "all":
            df = df[df[filter_col].astype(str) == str(filter_val)]

    if not x_axis or not y_axis:
        return jsonify({"error": "x_axis and y_axis are required"}), 400
    if x_axis not in df.columns:
        return jsonify({"error": "Invalid x_axis column"}), 400
    if y_axis not in df.columns and y_axis != "Count":
        return jsonify({"error": "Invalid y_axis column"}), 400

    work = df[[x_axis]].copy()
    if y_axis == "Count":
        work["__value__"] = 1
    else:
        work[y_axis] = pd.to_numeric(df[y_axis], errors="coerce").fillna(0)
        work["__value__"] = work[y_axis]

    if (work[x_axis].dtype == object or pd.api.types.is_string_dtype(work[x_axis])) and not pd.api.types.is_datetime64_any_dtype(work[x_axis]):
        parsed = pd.to_datetime(work[x_axis], errors="coerce")
        if parsed.notna().any():
            work[x_axis] = parsed

    grouped = work.groupby(x_axis, dropna=False)["__value__"].sum().reset_index()
    grouped = _sort_chart_labels(grouped, x_axis)

    labels = []
    if pd.api.types.is_datetime64_any_dtype(grouped[x_axis]):
        labels = grouped[x_axis].dt.strftime("%Y-%m-%d").tolist()
    else:
        labels = [str(v) for v in grouped[x_axis].tolist()]
    values = [float(v) for v in grouped["__value__"].tolist()]

    return jsonify({"labels": labels, "values": values})


if __name__ == "__main__":
    if SAMPLE_FILE.exists() and CURRENT_FILE is None:
        CURRENT_FILE = SAMPLE_FILE

    ip = _lan_ip()
    port = 5000
    print("\n" + "=" * 50)
    print("  Data Dashboard — LAN access")
    print("=" * 50)
    print(f"  Local:   http://127.0.0.1:{port}")
    print(f"  Network: http://{ip}:{port}")
    print("=" * 50 + "\n")

    app.run(host="0.0.0.0", port=port, debug=not IS_FROZEN, use_reloader=not IS_FROZEN)
